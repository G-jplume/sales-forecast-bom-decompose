# -*- coding: utf-8 -*-
"""
销售预测 捆绑SKU拆解 Streamlit 小程序
=====================================
自包含：decompose.py 与 app.py 在同一目录。
- Windows 本地运行：支持 DRM 加密文件（通过 Excel COM）
- Streamlit Cloud (Linux) 运行：自动降级为 openpyxl 读取（仅非加密文件）

运行: streamlit run app.py
"""
import sys
import os
import io
import math
import tempfile
import shutil
import traceback
from datetime import datetime
from collections import defaultdict

# ---- 页面配置（必须是第一个 Streamlit 命令）----
import streamlit as st

st.set_page_config(
    page_title="销售预测捆绑SKU拆解",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded",
)

import openpyxl

# ---- 导入 decompose 模块（同目录）----
APP_DIR = os.path.dirname(os.path.abspath(__file__))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

_DC_IMPORT_ERROR = None
try:
    import decompose as dc
except ImportError as e:
    _DC_IMPORT_ERROR = str(e)
    dc = None

# ---- 检测运行环境 ----
IS_WINDOWS = sys.platform == "win32"
HAS_PYWIN32 = False
if IS_WINDOWS:
    try:
        import win32com.client
        HAS_PYWIN32 = True
    except ImportError:
        pass


# ==================== 样式 ====================


# ==================== openpyxl 降级读取 ====================

def _read_sheet_openpyxl(wb_path, sheet_name, start_row=1):
    """用 openpyxl 读取 sheet（非 DRM 文件降级方案）。"""
    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=start_row, values_only=True))
    wb.close()
    result = []
    for r in rows:
        if isinstance(r, tuple):
            result.append(list(r))
        else:
            result.append([r])
    return result


def _detect_sheet_name(wb_path, keyword):
    """自动检测含关键字的 sheet 名。"""
    wb = openpyxl.load_workbook(wb_path, read_only=True)
    for name in wb.sheetnames:
        if keyword in name:
            wb.close()
            return name
    wb.close()
    return wb.sheetnames[0] if wb.sheetnames else None


def read_workbook_fallback(op_path, op_sheet=None, read_bom=True):
    """openpyxl 降级版 read_workbook（无 COM 环境）。"""
    op_name = op_sheet or _detect_sheet_name(op_path, "运营")
    bom_name = None

    wb = openpyxl.load_workbook(op_path, data_only=True, read_only=True)
    names = wb.sheetnames
    if read_bom:
        bom_name = next((n for n in names if "捆绑" in n), None)
    wb.close()

    op_rows = _read_sheet_openpyxl(op_path, op_name)

    def _norm(v):
        if v is None:
            return ""
        if hasattr(v, 'year'):
            return "%d年%d月" % (v.year, v.month)
        return v

    op_rows = [[_norm(v) for v in (r + [""] * (max(len(rr) for rr in op_rows) - len(r)))]
               if isinstance(r, list) else [_norm(r)] for r in op_rows]

    bom_rows = []
    if bom_name:
        bom_rows = _read_sheet_openpyxl(op_path, bom_name)
        bom_rows = [[_norm(v) for v in (r + [""] * (max(len(rr) for rr in bom_rows) - len(r)))]
                    if isinstance(r, list) else [_norm(r)] for r in bom_rows]

    return names, op_name, bom_name, op_rows, bom_rows


def read_normalization_fallback(inv_path):
    """openpyxl 降级版 read_normalization。"""
    shop_to_main = {}
    person_to_norm = {}

    ws_name = _detect_sheet_name(inv_path, "海外仓库存")
    if not ws_name:
        return shop_to_main, person_to_norm

    data = _read_sheet_openpyxl(inv_path, ws_name, start_row=2)
    for row in data:
        if not row or not row[0]:
            continue
        main_shop = dc.safe_str(row[2]) if len(row) > 2 else ""
        norm_person = dc.safe_str(row[4]) if len(row) > 4 else ""
        shop = dc.safe_str(row[5]) if len(row) > 5 else ""
        person = dc.safe_str(row[6]) if len(row) > 6 else ""
        if shop and main_shop:
            shop_to_main[shop] = main_shop
        if person and norm_person:
            person_to_norm[person] = norm_person

    return shop_to_main, person_to_norm


def read_inventory_fallback(inv_path):
    """openpyxl 降级版 read_inventory。"""
    inv_map = defaultdict(lambda: dict(zip(dc.INV_COMPONENTS, [0]*7)))

    sheet_map = {
        "海外仓库存": ("海外在仓", 7),
        "在途&配柜": ("海外在途", 12),
        "合同待交付": ("待交付合同", None),
        "采购单待下单": ("国内待下合同", 14),
    }

    for sheet_name, (comp, qty_col) in sheet_map.items():
        try:
            data = _read_sheet_openpyxl(inv_path, sheet_name, start_row=2)
        except Exception:
            continue
        for row in data:
            if not row or not row[0]:
                continue
            key = dc.safe_str(row[0])
            if sheet_name == "合同待交付":
                qty = (dc.to_num(row[19]) if len(row) > 19 else 0) + \
                      (dc.to_num(row[25]) if len(row) > 25 else 0)
            else:
                qty = dc.to_num(row[qty_col]) if len(row) > qty_col else 0
            inv_map[key][comp] += qty

    # 海外仓调拨
    try:
        data = _read_sheet_openpyxl(inv_path, "海外仓调拨", start_row=3)
        for row in data:
            if not row:
                continue
            key_in = dc.safe_str(row[0]) if len(row) > 0 else ""
            out_sku = dc.safe_str(row[2]) if len(row) > 2 else ""
            out_shop = dc.safe_str(row[6]) if len(row) > 6 else ""
            out_person = dc.safe_str(row[8]) if len(row) > 8 else ""
            out_mode = dc.safe_str(row[9]) if len(row) > 9 else ""
            qty = dc.to_num(row[14]) if len(row) > 14 else 0
            if key_in:
                inv_map[key_in]["借入"] += qty
            if out_sku and out_shop and out_mode and out_person:
                key_out = "%s&%s&%s&%s" % (out_sku, out_shop, out_mode, out_person)
                inv_map[key_out]["借出"] += qty
    except Exception:
        pass

    return dict(inv_map)


# ==================== 飞书映射降级 ====================

def read_feishu_mapping_safe(feishu_url):
    """尝试读取飞书映射，失败时返回空映射。"""
    try:
        return dc.read_feishu_mapping(feishu_url)
    except FileNotFoundError:
        st.warning("⚠️ lark-cli 未安装或未授权，跳过飞书映射。Sheet1/整套替代功能不可用。")
        return {}, {}, {}, {}
    except Exception as e:
        st.warning(f"⚠️ 飞书映射读取失败: {e}。Sheet1/整套替代功能不可用。")
        return {}, {}, {}, {}


# ==================== 核心拆解流程 ====================

def run_decompose(op_path, bom_csv_path, product_csv_path, inv_path, feishu_url,
                  op_sheet=None, enc=None):
    """执行拆解流水线，返回 (output_path, summary_dict)。"""
    summary = {}

    # ---- Step 1: 读取运营表 ----
    st.write("📊 读取运营表...")

    def _try_read_workbook():
        if HAS_PYWIN32:
            return dc.read_workbook(op_path, op_sheet=op_sheet, read_bom=True)
        else:
            return read_workbook_fallback(op_path, op_sheet=op_sheet, read_bom=True)

    try:
        names, op_name, bom_name, op, bom = _try_read_workbook()
    except Exception as e:
        if not HAS_PYWIN32:
            st.error(
                f"读取运营表失败: {e}\n\n"
                "可能原因：文件被 DRM 加密，当前环境（Linux/Cloud）无法读取。\n"
                "解决方案：在 Windows + Excel 环境下本地运行 `streamlit run app.py`。"
            )
            raise
        raise

    st.write(f"  运营表: `{op_name}`")

    # ---- Step 2: 定位表头 / 月份行 / 数据行 ----
    st.write("📋 定位表头和数据...")
    hidx = dc.find_header_row(op, "SKU", col=0)
    if hidx == 0 and not str(dc.cell(op[0], 0)).strip() == "SKU":
        st.warning("未在运营表定位到表头行，已回退到第 1 行")
    op_hdr = op[hidx]
    month_row = op[hidx + 1] if hidx + 1 < len(op) else []
    data = [r for r in op[hidx + 2:] if r and str(dc.cell(r, 0)).strip()]
    st.write(f"  数据行数: {len(data)}")

    # ---- Step 3: 列索引 ----
    fstart = dc.col_index(op_hdr, "销量预测")
    if fstart is None:
        fstart = dc.DEFAULT_FSTART
    ss_col = fstart - 1
    fc = list(range(fstart, fstart + 12))
    cc = {
        "fstart": fstart, "ss_col": ss_col, "tot_col": fstart - 2, "fc": fc,
        "sku_c": dc.col_index(op_hdr, "SKU", exclude="产品") or 0,
        "shop_c": dc.col_index(op_hdr, "店铺"),
        "opg_c": dc.col_index(op_hdr, "运营"),
        "ship_c": dc.col_index(op_hdr, "发货模式"),
        "cat_c": dc.col_index(op_hdr, "一级分类"),
        "grade_c": dc.col_index(op_hdr, "产品定级"),
        "chan_c": dc.col_index(op_hdr, "渠道"),
        "coun_c": dc.col_index(op_hdr, "国家"),
        "stat_c": dc.col_index(op_hdr, "销售状态"),
        "list_c": dc.col_index(op_hdr, "listing") or dc.col_index(op_hdr, "负责人"),
        "msku_c": dc.col_index(op_hdr, "MSKU"),
        "pname_c": dc.col_index(op_hdr, "品名描述"),
        "spu_c": dc.col_index(op_hdr, "型号SPU"),
    }

    # ---- Step 4: 属性查找表 ----
    op_attr = {}
    for row in data:
        sku = str(dc.cell(row, cc["sku_c"])).strip()
        shop = str(dc.cell(row, cc["shop_c"])).strip()
        opg = str(dc.cell(row, cc["opg_c"])).strip()
        ship = str(dc.cell(row, cc["ship_c"])).strip()
        vals = (
            str(dc.cell(row, cc["cat_c"])).strip(),
            str(dc.cell(row, cc["grade_c"])).strip(),
            str(dc.cell(row, cc["chan_c"])).strip(),
            str(dc.cell(row, cc["coun_c"])).strip(),
            str(dc.cell(row, cc["stat_c"])).strip(),
        )
        key = (sku.upper(), shop, opg, ship)
        if key not in op_attr:
            op_attr[key] = vals

    # ---- Step 5: 构建 BOM / 产品资料 ----
    st.write("🔧 构建BOM和产品资料...")
    if bom_csv_path and product_csv_path:
        prod_master, bom_map = dc.build_bom_product(bom_csv_path, product_csv_path, enc)
        st.write(f"  产品资料: {len(prod_master)} 条 | BOM映射: {len(bom_map)} 条")
    else:
        prod_master, bom_map = {}, defaultdict(list)
        st.warning("未提供 BOM CSV 和产品资料 CSV")

    # ---- Step 6: 构建明细 ----
    st.write("📦 构建拆解明细...")
    (detail, missing_bundle, missing_prod,
     n_bundle, n_single, n_missbundle, rounded_cells) = \
        dc.process_rows(data, bom_map, prod_master, op_attr, cc)
    st.write(f"  明细总行数: **{len(detail)}** (捆绑拆解: {n_bundle}, 原单品: {n_single})")

    if missing_bundle:
        st.write(f"  ⚠️ 缺失捆绑关系: {len(missing_bundle)} 个 → {', '.join(m[0] for m in missing_bundle)}")
    if missing_prod:
        st.write(f"  ⚠️ 未匹配产品资料: {len(missing_prod)} 个")

    summary["明细总行数"] = len(detail)
    summary["捆绑拆解"] = n_bundle
    summary["原单品"] = n_single
    summary["缺失捆绑关系"] = n_missbundle
    summary["未匹配产品资料"] = len(missing_prod)

    # ---- Step 7: 自检 ----
    formula_bad = dc.check_formulas(detail)
    if formula_bad:
        st.warning(f"  汇总公式不符: {len(formula_bad)} 行")
    else:
        st.write("  ✅ 汇总公式自检通过")

    # ---- Step 8: Sheet1/Sheet2 生成（需要库存文件） ----
    sheet1_data = None
    sheet2_data = None
    months = None

    if inv_path:
        try:
            st.write("🔗 读取飞书映射表...")
            src_to_map, zhu_to_sets, zhu_to_xiangsi, zhu_to_order = \
                read_feishu_mapping_safe(feishu_url)

            if not src_to_map:
                st.info("飞书映射为空，跳过 Sheet1/整套替代生成。"
                        "仅输出 Sheet2(未拆解) + Sheet3(异常提醒)。")
            else:
                st.write(f"  源SKU映射: {len(src_to_map)} | 整套替代: {len(zhu_to_sets)}")

                st.write("📐 读取归一化映射...")
                if HAS_PYWIN32:
                    shop_to_main, person_to_norm = dc.read_normalization(inv_path)
                else:
                    shop_to_main, person_to_norm = read_normalization_fallback(inv_path)
                st.write(f"  店铺→主店铺: {len(shop_to_main)} | 运营→归一: {len(person_to_norm)}")

                st.write("🔄 转换明细为records...")
                records, months = dc.detail_to_records(detail, month_row, fstart)
                st.write(f"  records: {len(records)} | months: {len(months)}")

                st.write("📊 读取库存文件(5个sheet)...")
                if HAS_PYWIN32:
                    inv_map = dc.read_inventory(inv_path)
                else:
                    inv_map = read_inventory_fallback(inv_path)
                st.write(f"  库存key: {len(inv_map)}")

                st.write("📦 生成捆绑拆分...")
                sheet2_data = dc.build_sheet2_bundle(
                    records, months, src_to_map, zhu_to_sets,
                    zhu_to_xiangsi, zhu_to_order, inv_map,
                    shop_to_main, person_to_norm
                )
                zhengtao = sum(1 for r in sheet2_data if r["is_zhengtao"])
                results_ok = sum(1 for r in sheet2_data if r["结果"] == "Y")
                matched = sum(1 for r in sheet2_data if any(r["set_total_inv"]))
                st.write(f"  Sheet2行数: {len(sheet2_data)} | 整套替代: {zhengtao} | "
                         f"结果Y: {results_ok} N: {len(sheet2_data)-results_ok} | 有库存: {matched}")

                summary["Sheet2行数"] = len(sheet2_data)
                summary["整套替代"] = zhengtao
                summary["结果Y"] = results_ok
                summary["有库存"] = matched

                st.write("📋 生成销售需求汇总（含按月从老到新分配）...")
                sheet1_data = dc.build_sheet1(
                    records, months, src_to_map, zhu_to_sets,
                    zhu_to_order, shop_to_main, person_to_norm,
                    sheet2_data
                )
                st.write(f"  Sheet1行数: **{len(sheet1_data)}**")

                summary["Sheet1行数"] = len(sheet1_data)
        except Exception as e:
            st.error(f"Sheet1/Sheet2 生成失败: {e}")
            traceback.print_exc()
            sheet1_data = None
            sheet2_data = None
            months = None

    # ---- Step 9: 输出文件 ----
    st.write("💾 生成输出文件...")
    has_sheet1 = sheet1_data is not None and months
    if has_sheet1:
        out_name = "销售需求汇总_整套分配_%s.xlsx" % datetime.now().strftime("%Y%m%d")
    else:
        out_name = "销售预测拆解汇总_明细_%s.xlsx" % datetime.now().strftime("%Y%m%d")
    out_path = os.path.join(tempfile.gettempdir(), out_name)

    sheetnames = dc.write_output(
        out_path, detail, op_hdr, month_row, fstart,
        missing_bundle, missing_prod, formula_bad,
        sheet1_data=sheet1_data, months=months,
        sheet2_data=None,
    )
    st.write(f"  输出文件: `{out_name}`")
    st.write(f"  工作表: {', '.join(sheetnames)}")

    summary["工作表"] = sheetnames
    summary["汇总公式自检"] = "全部通过" if not formula_bad else f"{len(formula_bad)}行不符"

    return out_path, summary


# ==================== 主界面 ====================

def main():
    if _DC_IMPORT_ERROR:
        st.error(f"无法导入 decompose 模块: {_DC_IMPORT_ERROR}")
        st.stop()

    st.title("销售预测捆绑SKU拆解")
    st.caption("上传运营表 + BOM + 产品资料（+ 可选库存文件），一键拆解捆绑SKU并输出完整明细")

    # 环境提示
    if not HAS_PYWIN32:
        st.info("ℹ️ 当前运行环境无 Excel COM（非 Windows 或未安装 pywin32）。"
                "DRM 加密文件无法读取，请上传未加密的 xlsx 文件。"
                "如需处理 DRM 文件，请在 Windows + Excel 环境下本地运行。")

    # ---- 侧边栏 ----
    with st.sidebar:
        st.header("⚙️ 参数设置")
        feishu_url = st.text_input(
            "飞书整套替代关系表URL",
            value=dc.FEISHU_URL_DEFAULT,
            help="用于读取整套替代关系和混用SKU映射。需要 lark-cli 已授权。"
        )

        op_sheet_name = st.text_input(
            "运营表 Sheet名（留空自动识别）",
            value="",
            help="自动识别含「运营」的sheet"
        )

        csv_enc = st.selectbox(
            "CSV编码",
            options=["自动探测", "gbk", "utf-8-sig", "utf-8", "latin-1"],
            help="默认自动探测，如遇乱码可手动指定"
        )

        st.divider()
        st.markdown("### 📖 使用说明")
        st.markdown(
            "1. 上传**运营表** xlsx（支持DRM加密）\n"
            "2. 上传**同款套装库存** CSV + **产品资料** CSV（成对必填）\n"
            "3. 可选：上传**库存文件** xlsx → 额外生成Sheet1「销售需求汇总」\n"
            "4. 点击「开始拆解」\n"
            "5. 完成后点击「下载结果文件」"
        )
        st.divider()
        env_label = "Windows + Excel COM" if HAS_PYWIN32 else "Linux/Cloud (openpyxl)"
        st.caption(f"🖥️ 运行环境: {env_label}")
        st.caption("📦 自包含 decompose.py 模块")

    # ---- 文件上传区 ----
    st.markdown("### 📁 文件上传")

    col1, col2 = st.columns(2)

    with col1:
        st.markdown("**必填文件**")
        op_file = st.file_uploader(
            "运营表 xlsx",
            type=["xlsx"],
            key="op_file",
            help="运营-销售预测汇总表（DRM加密需在Windows本地运行）"
        )
        bom_csv = st.file_uploader(
            "同款套装库存 CSV",
            type=["csv"],
            key="bom_csv",
            help="包含 仓库SKU / 父仓库SKU / 套装数量"
        )

    with col2:
        st.markdown("**必填文件**")
        product_csv = st.file_uploader(
            "产品资料 CSV",
            type=["csv"],
            key="product_csv",
            help="包含 产品SKU / 产品名称 / 产品款式"
        )
        st.markdown("**可选文件**")
        inv_file = st.file_uploader(
            "库存文件 xlsx（生成Sheet1）",
            type=["xlsx"],
            key="inv_file",
            help="含海外仓库存/在途配柜/合同待交付等5个sheet"
        )

    # ---- 运行按钮 ----
    st.divider()

    can_run = op_file is not None and bom_csv is not None and product_csv is not None

    if not can_run:
        st.info("👆 请上传运营表 + 同款套装库存CSV + 产品资料CSV 后开始拆解")

    if st.button("🚀 开始拆解", type="primary", disabled=not can_run):
        # ---- 保存上传文件到临时目录 ----
        with tempfile.TemporaryDirectory(prefix="decompose_") as tmpdir:
            op_path = os.path.join(tmpdir, op_file.name)
            with open(op_path, "wb") as f:
                shutil.copyfileobj(op_file, f)

            bom_path = os.path.join(tmpdir, bom_csv.name)
            with open(bom_path, "wb") as f:
                shutil.copyfileobj(bom_csv, f)

            prod_path = os.path.join(tmpdir, product_csv.name)
            with open(prod_path, "wb") as f:
                shutil.copyfileobj(product_csv, f)

            inv_path = None
            if inv_file:
                inv_path = os.path.join(tmpdir, inv_file.name)
                with open(inv_path, "wb") as f:
                    shutil.copyfileobj(inv_file, f)

            # ---- 执行拆解 ----
            enc_val = None if csv_enc == "自动探测" else csv_enc
            sheet_val = op_sheet_name.strip() or None

            with st.status("拆解进行中...", expanded=True) as status:
                try:
                    result_path, summary = run_decompose(
                        op_path, bom_path, prod_path,
                        inv_path if inv_path else None,
                        feishu_url,
                        op_sheet=sheet_val,
                        enc=enc_val,
                    )
                    status.update(label="✅ 拆解完成！", state="complete")
                except Exception as e:
                    status.update(label=f"❌ 拆解失败: {e}", state="error")
                    st.error(f"拆解过程中出错: {e}")
                    st.code(traceback.format_exc(), language="python")
                    return

            # ---- 结果展示 ----
            st.divider()
            st.markdown("### 📊 运行摘要")

            cols = st.columns(min(len(summary), 6))
            metric_map = {
                "明细总行数": "总行数",
                "捆绑拆解": "捆绑拆解",
                "原单品": "原单品",
                "Sheet1行数": "Sheet1行数",
                "Sheet2行数": "Sheet2行数",
                "有库存": "有库存行",
                "结果Y": "结果Y",
                "整套替代": "整套替代",
                "缺失捆绑关系": "缺失关系",
                "未匹配产品资料": "未匹配资料",
            }
            for i, (k, v) in enumerate(summary.items()):
                if k in ("工作表", "汇总公式自检"):
                    continue
                if i >= len(cols):
                    break
                label = metric_map.get(k, k)
                cols[i].metric(label, v)

            if "工作表" in summary:
                st.write(f"**工作表**: {', '.join(summary['工作表'])}")
            if "汇总公式自检" in summary:
                st.write(f"**汇总公式自检**: {summary['汇总公式自检']}")

            # ---- 下载按钮 ----
            st.divider()
            st.markdown("### 📥 下载结果")

            with open(result_path, "rb") as f:
                file_data = f.read()

            out_filename = os.path.basename(result_path)
            st.download_button(
                label="📥 下载结果文件",
                data=file_data,
                file_name=out_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

            st.success(f"文件已生成: `{out_filename}` ({len(file_data):,} bytes)")


if __name__ == "__main__":
    main()
